"""Small research display of existing H1 metrics; never a second TR calculator.

The producer supplies sealed H1 results, issuer metadata and dated holdings.
This module performs cross-sectional ranking only. No network, source admission,
portfolio selection, scheduled collection or accepted-session publication.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import os
import stat
import zipfile
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlsplit

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from tools.etf116_eod_gate import EXPECTED_UNIVERSE_SHA256, validate_universe
from tools.etf116_research_metrics_h1 import canonical_hash, d, ts

SCHEMA = 'etf116-simple-report-input-v1'
OUTPUT_SCHEMA = 'etf116-simple-report-v1'
HORIZONS = ('12', '36', '60')
WEIGHTS = {'cagr_1y': .20, 'cagr_3y': .20, 'cagr_5y': .20,
           'sharpe_3y': .15, 'mdd_3y': .10, 'volatility_3y': .05,
           'expense_ratio': .10}
SHA = re.compile(r'[a-f0-9]{64}')
SYMBOL = re.compile(r'[A-Z][A-Z0-9.\-]{0,15}')
FLAGS = ('leveraged', 'inverse', 'option_overlay', 'single_stock', 'fund_of_funds')
CURRENCIES = frozenset('USD EUR GBP JPY CHF CAD AUD NZD HKD CNY CNH KRW SGD TWD INR BRL MXN ZAR SEK NOK DKK PLN ILS AED SAR'.split())


class ReportRefused(ValueError):
    pass


def require(ok, code):
    if not ok:
        raise ReportRefused(code)


def number(value, code):
    require(type(value) in (int, float) and math.isfinite(value), code)
    return float(value)


def source(record, *, session, decision, max_age):
    """Only explicit plain provenance, never credential URLs or error text."""
    url = record.get('source_url')
    require(type(url) is str and len(url) <= 512 and not any(c.isspace() for c in url)
            and '\\' not in url, 'SOURCE_URL')
    try:
        parts = urlsplit(url)
        require(parts.scheme == 'https' and bool(parts.hostname) and not parts.username
                and not parts.password and not parts.query and not parts.fragment
                and parts.port in (None, 443), 'SOURCE_URL')
    except ValueError:
        raise ReportRefused('SOURCE_URL') from None
    require(type(record.get('source_sha256')) is str
            and SHA.fullmatch(record['source_sha256']), 'SOURCE_HASH')
    observed, released = d(record['asof']), ts(record['published_at'])
    require(0 <= (d(session) - observed).days <= max_age
            and observed <= released.date() and released <= decision, 'SOURCE_CLOCK')


def percentile(values, value):
    """Average ranks for ties; a constant cohort contributes a neutral 50%."""
    if len(values) == 1:
        return .5
    return (sum(x < value for x in values) + (sum(x == value for x in values) - 1) / 2) / (len(values) - 1)


def _select(packet, frozen, *, evaluation_time=None):
    """Shared selection before holdings I/O; never ranks a truncated cohort."""
    now = evaluation_time or datetime.now(timezone.utc)
    require(now.tzinfo is not None, 'EVALUATION_CLOCK')
    require(type(packet) is dict and packet.get('schema') == SCHEMA, 'INPUT_SCHEMA')
    tickers = validate_universe(frozen)
    metrics = packet['metrics_candidate']
    require(type(metrics) is dict and metrics.get('schema') == 'etf116-metrics-candidate-output-v1'
            and metrics.get('universe_sha256') == EXPECTED_UNIVERSE_SHA256
            and metrics.get('research_status') == 'PREPARED_ONLY_NO_AUTHORITY'
            and metrics.get('economic_authority') is False
            and metrics.get('allowed_for_sheet_publisher') is False, 'METRICS_CONTRACT')
    sealed = {k: v for k, v in metrics.items() if k != 'candidate_receipt_sha256'}
    require(metrics.get('candidate_receipt_sha256') == canonical_hash(sealed), 'METRICS_SEAL')
    session = metrics['session']
    decision = ts(packet['decision_at'])
    require(d(session) <= decision.date() and decision <= now
            and timedelta(0) <= now - decision <= timedelta(days=4)
            and 0 <= (decision.date() - d(session)).days <= 4, 'REPORT_CLOCK')
    require(type(metrics.get('source_input_bytes_sha256')) is str
            and SHA.fullmatch(metrics['source_input_bytes_sha256']), 'METRICS_SOURCE_HASH')
    metadata, results = packet['fund_metadata'], metrics['ticker_results']
    require(type(metadata) is dict and set(metadata) == set(tickers)
            and type(results) is dict and set(results) == set(tickers), 'FIXED_UNIVERSE_COVERAGE')
    rows, excluded, anchors, basis = [], {}, None, None
    for ticker in tickers:
        meta = metadata[ticker]
        require(type(meta) is dict and meta.get('asset_class') in
                ('EQUITY', 'BOND', 'COMMODITY', 'MIXED', 'OTHER'), 'FUND_CLASSIFICATION')
        require(all(type(meta.get(k)) is bool for k in FLAGS), 'FUND_CLASSIFICATION')
        require(type(meta.get('currency')) is str and meta['currency'] in CURRENCIES, 'FUND_CURRENCY')
        source(meta, session=session, decision=decision, max_age=45)
        if meta['asset_class'] != 'EQUITY' or any(meta[k] for k in FLAGS) or meta.get('currency') != 'USD':
            excluded[ticker] = 'OUTSIDE_GENERAL_USD_EQUITY_COHORT'
            continue
        report = results[ticker]
        windows = report.get('horizons', {})
        if any(windows.get(h, {}).get('status') != 'RESEARCH_DAILY_TR_COMPLETE' for h in HORIZONS):
            excluded[ticker] = 'INCOMPLETE_1_3_5_YEAR_HISTORY'
            continue
        require(report.get('ticker') == ticker and report.get('asof') == session
                and report.get('series_basis') in ('NAV_TR_INDEX', 'MARKET_TR_INDEX'), 'METRICS_IDENTITY')
        risk = windows['36'].get('risk', {})
        if risk.get('risk_free_status') != 'RESEARCH_RF_AVAILABLE' or risk.get('sharpe') is None:
            excluded[ticker] = 'DATED_RF_SHARPE_UNAVAILABLE'
            continue
        expense = number(meta.get('expense_ratio'), 'EXPENSE_RATIO')
        require(0 <= expense <= .05, 'EXPENSE_RATIO_UNIT')
        vol = number(risk.get('ann_volatility'), 'VOLATILITY')
        mdd = number(risk.get('mdd'), 'MDD')
        require(vol >= 0 and -1 <= mdd <= 0, 'RISK_UNIT')
        name = meta.get('fund_name')
        require(type(name) is str and 0 < len(name) <= 160
                and all(c.isprintable() for c in name), 'FUND_NAME')
        require(all(number(windows[h].get('cagr'), 'CAGR') > -1 for h in HORIZONS), 'CAGR_UNIT')
        # Excluded products/missing history/RF cannot establish a cohort basis.
        these_anchors = tuple(windows[h]['anchor_session'] for h in HORIZONS)
        if anchors is None:
            anchors, basis = these_anchors, report['series_basis']
        require(these_anchors == anchors and report['series_basis'] == basis
                and all(windows[h].get('basis') == basis for h in HORIZONS), 'NONCOMPARABLE_WINDOWS_OR_BASIS')
        rows.append(dict(ticker=ticker, name=name,
                         cagr_1y=number(windows['12'].get('cagr'), 'CAGR'),
                         cagr_3y=number(windows['36'].get('cagr'), 'CAGR'),
                         cagr_5y=number(windows['60'].get('cagr'), 'CAGR'),
                         sharpe_3y=number(risk.get('sharpe'), 'SHARPE'), mdd_3y=mdd,
                         volatility_3y=vol, expense_ratio=expense,
                         metadata_asof=meta['asof'], metadata_source=meta['source_url']))
    require(len(rows) >= 5, 'FEWER_THAN_FIVE_QUALIFIED_EQUITY_ETFS')
    for row in rows:
        components = {}
        for key, weight in WEIGHTS.items():
            sign = -1 if key in ('volatility_3y', 'expense_ratio') else 1
            components[key] = percentile([sign * r[key] for r in rows], sign * row[key])
        row['score'] = 100 * sum(WEIGHTS[key] * components[key] for key in WEIGHTS)
        row['percentiles'] = components
    rows.sort(key=lambda r: (-r['score'], -r['cagr_3y'], r['expense_ratio'], r['ticker']))
    top = [dict(row, rank=i + 1) for i, row in enumerate(rows[:5])]
    return dict(session=session, decision=decision, tickers=tickers, metrics=metrics,
                rows=rows, excluded=excluded, basis=basis, top=top)


def _holdings(top, supplied, *, session, decision, tickers):
    holdings, overlap = [], {}
    require(type(supplied) is dict and set(supplied).issubset(set(tickers)), 'HOLDINGS_SCOPE')
    for fund in top:
        ticker = fund['ticker']
        records = supplied.get(ticker)
        require(type(records) is list and len(records) == 10, 'TOP10_HOLDINGS_REQUIRED')
        require(all(type(r) is dict for r in records), 'HOLDINGS_RECORDS')
        require({r.get('rank') for r in records} == set(range(1, 11))
                and all(type(r.get('rank')) is int for r in records), 'HOLDINGS_RANKS')
        seen, total, snapshot = set(), 0., set()
        for rec in sorted(records, key=lambda r: r['rank']):
            source(rec, session=session, decision=decision, max_age=10)
            symbol, company = rec.get('symbol'), rec.get('company_name')
            require(type(symbol) is str and SYMBOL.fullmatch(symbol) and symbol not in seen
                    and rec.get('asset_class') == 'EQUITY', 'STOCK_HOLDING_IDENTITY')
            require(type(company) is str and 0 < len(company) <= 160
                    and all(c.isprintable() for c in company), 'COMPANY_NAME')
            seen.add(symbol)
            weight = number(rec.get('weight_pct'), 'HOLDING_WEIGHT')
            require(0 < weight <= 100, 'HOLDING_WEIGHT_UNIT')
            total += weight
            snapshot.add((rec['asof'], rec['source_url'], ts(rec['published_at']), rec['source_sha256']))
            holdings.append(dict(etf_rank=fund['rank'], etf=ticker, rank=rec['rank'],
                                 symbol=symbol, company_name=company, weight_pct=weight,
                                 asof=rec['asof'], source_url=rec['source_url']))
            member = overlap.setdefault(symbol, dict(symbol=symbol, etfs=[], fund_count=0))
            member['etfs'].append(ticker)
            member['fund_count'] += 1
        require(total <= 100 + 1e-9 and len(snapshot) == 1, 'HOLDINGS_SNAPSHOT')
        require(all(a['weight_pct'] >= b['weight_pct'] for a, b in zip(
            sorted(records, key=lambda r: r['rank']), sorted(records, key=lambda r: r['rank'])[1:])), 'HOLDINGS_WEIGHT_ORDER')
    return holdings, overlap


def plan_inputs(packet, frozen, *, evaluation_time=None, holdings_refresh_due=False):
    """Plan GET-only intake for selected funds. Does not execute or admit inputs.

    The existing cadence owner supplies refresh_due; no scheduler or new clock
    policy is introduced. Every later export revalidates the original inputs.
    """
    require(type(holdings_refresh_due) is bool, 'HOLDINGS_REFRESH_FLAG')
    selected = _select(packet, frozen, evaluation_time=evaluation_time)
    cached = packet.get('holdings', {})
    require(type(cached) is dict and set(cached).issubset(set(selected['tickers'])), 'HOLDINGS_SCOPE')
    fetch, reuse = {}, {}
    for fund in selected['top']:
        ticker = fund['ticker']
        if holdings_refresh_due:
            fetch[ticker] = 'EXISTING_CADENCE_REFRESH_DUE'
        elif ticker not in cached:
            fetch[ticker] = 'MISSING_HOLDINGS_CACHE'
        else:
            try:
                _holdings([fund], {ticker: cached[ticker]}, session=selected['session'],
                          decision=selected['decision'], tickers=selected['tickers'])
                reuse[ticker] = canonical_hash(cached[ticker])
            except (ValueError, KeyError, TypeError, OverflowError):
                # Cache invalidity requests fresh GETs, never output/error echo.
                fetch[ticker] = 'INVALID_OR_STALE_HOLDINGS_CACHE'
    result = dict(schema='etf116-simple-report-read-plan-v1', session=selected['session'],
                  input_sha256=canonical_hash(packet),
                  metrics_receipt_sha256=selected['metrics']['candidate_receipt_sha256'],
                  selected_tickers=[r['ticker'] for r in selected['top']],
                  qualified_count=len(selected['rows']), frozen_count=len(selected['tickers']),
                  holdings_fetch_reasons=fetch, holdings_reuse_sha256=reuse,
                  funds_requiring_holdings=len(fetch), report_input_ready=not fetch,
                  network_calls_performed=0, get_only=True, metrics_recompute_requested=False,
                  publish_enabled=False, research_only=True, economic_authority=False,
                  source_admission=False, accepted_session=False)
    return dict(result, plan_sha256=canonical_hash(result))


def build(packet, frozen, *, evaluation_time=None):
    selected = _select(packet, frozen, evaluation_time=evaluation_time)
    session, metrics, top = selected['session'], selected['metrics'], selected['top']
    holdings, overlap = _holdings(top, packet.get('holdings'), session=session,
                                 decision=selected['decision'], tickers=selected['tickers'])
    report = dict(schema=OUTPUT_SCHEMA, session=session, decision_at=packet['decision_at'],
                  research_only=True, economic_authority=False, accepted_session=False,
                  allowed_for_existing_sheet_publisher=False, source_admission=False,
                  input_sha256=canonical_hash(packet), metrics_receipt_sha256=metrics['candidate_receipt_sha256'],
                  series_basis=selected['basis'], qualified_count=len(selected['rows']), frozen_count=len(selected['tickers']),
                  weights=WEIGHTS, top5=top, holdings=holdings, excluded=selected['excluded'],
                  overlap=sorted(overlap.values(), key=lambda r: (-r['fund_count'], r['symbol'])),
                  method='Equal 1/3/5Y CAGR percentile weights20% each; 3Y risk; fee preference10%.',
                  fee_note='Fund expenses are already reflected in TR; no second return deduction.',
                  scope_note='Top5 among qualified general USD equity ETFs in frozen116, not the whole ETF market.')
    return dict(report, report_sha256=canonical_hash(report))


def workbook_bytes(report):
    """Public Actions runtime dependency; values only, no executable formulas."""
    # The desktop artifact authoring runtime is not distributed to GitHub runners.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    def tab(name, title, headers, records, keys, percents=()):
        sheet = book.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.append([title])
        sheet.append(['As of', d(report['session']), 'Qualified ETFs', report['qualified_count'], 'Frozen universe', report['frozen_count']])
        sheet.append(['RESEARCH ONLY', report['scope_note']])
        sheet.append(headers)
        for row in records:
            sheet.append([d(row[k]) if k in ('asof', 'metadata_asof') else row[k] for k in keys])
        # Force external text to literal strings, including names beginning with '='.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top')
        for rownum in (1, 4):
            for cell in sheet[rownum]:
                cell.font = Font(color='FFFFFF', bold=True, size=13 if rownum == 1 else 10)
                cell.fill = PatternFill('solid', fgColor='17365D')
        sheet.freeze_panes = 'C5'
        sheet['B2'].number_format = 'yyyy-mm-dd'
        sheet.auto_filter.ref = f'A4:{get_column_letter(len(headers))}{4 + len(records)}'
        for i, key in enumerate(keys, 1):
            sheet.column_dimensions[get_column_letter(i)].width = 46 if key in ('name', 'company_name', 'source_url', 'metadata_source') else 19
            for rownum in range(5, 5 + len(records)):
                cell = sheet.cell(rownum, i)
                if key in percents:
                    cell.number_format = '0.00%;[Red](0.00%);"–"'
                elif key in ('score', 'sharpe_3y', 'weight_pct'):
                    cell.number_format = '0.00'
                elif key in ('asof', 'metadata_asof'):
                    cell.number_format = 'yyyy-mm-dd'
        # The overlap table has three columns, its context row has six.
        for index in range(len(keys) + 1, 7):
            sheet.column_dimensions[get_column_letter(index)].width = 19
        if name == 'OVERLAP':
            sheet.column_dimensions['C'].width = 46
        return sheet
    keys = ('rank', 'ticker', 'name', 'score', 'cagr_1y', 'cagr_3y', 'cagr_5y', 'sharpe_3y', 'mdd_3y', 'volatility_3y', 'expense_ratio', 'metadata_asof', 'metadata_source')
    s = tab('ETF TOP5', 'ETF Top5 · 1 / 3 / 5 years',
            ['Rank', 'ETF', 'Fund', 'Score /100', 'CAGR 1Y', 'CAGR 3Y', 'CAGR 5Y', 'Sharpe 3Y', 'MDD 3Y', 'Volatility 3Y', 'Annual fee', 'Fee as of', 'Issuer source'],
            report['top5'], keys, ('cagr_1y', 'cagr_3y', 'cagr_5y', 'mdd_3y', 'volatility_3y', 'expense_ratio'))
    s.append([])
    s.append(['Method', report['method']])
    s.append(['Expenses', report['fee_note']])
    s.append(['Basis', report['series_basis']])
    s.append(['Input SHA256', report['input_sha256']])
    s.append(['Report SHA256', report['report_sha256']])
    tab('TOP10 HOLDINGS', 'Top10 stocks in each selected ETF · issuer snapshot dates differ',
        ['ETF rank', 'ETF', 'Holding rank', 'Stock / class', 'Company', 'Fund weight (%)', 'Holdings as of', 'Issuer source'],
        report['holdings'], ('etf_rank', 'etf', 'rank', 'symbol', 'company_name', 'weight_pct', 'asof', 'source_url'))
    members = [dict(r, etfs=', '.join(r['etfs'])) for r in report['overlap']]
    tab('OVERLAP', 'Stock overlap · counts only; no assumed portfolio weights',
        ['Stock / class', 'ETF count', 'ETFs'], members, ('symbol', 'fund_count', 'etfs'))
    output = io.BytesIO()
    book.save(output)
    # Deterministic bytes make cache verification independent of self-rehashed
    # sidecar receipts. Document timestamps describe this input, not a new run.
    stable = io.BytesIO()
    with zipfile.ZipFile(output) as source_zip, zipfile.ZipFile(stable, 'w', compression=zipfile.ZIP_DEFLATED) as dest:
        for name in sorted(source_zip.namelist()):
            raw = source_zip.read(name)
            if name == 'docProps/core.xml':
                timestamp = ts(report['decision_at']).strftime('%Y-%m-%dT%H:%M:%SZ').encode()
                # Preserve namespace/QName declarations, including xsi:type.
                pattern = rb'(<dcterms:(created|modified)\b[^>]*>)[^<]*(</dcterms:\2>)'
                raw, count = re.subn(pattern, lambda match: match[1] + timestamp + match[3], raw)
                require(count == 2, 'WORKBOOK_CORE_TIMESTAMPS')
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            dest.writestr(info, raw)
    return stable.getvalue()


def publish_local(report, folder):
    """Content-addressed, exclusive local export. No overwrite or network retry."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / (report['session'] + '-' + report['report_sha256'])
    xlsx = workbook_bytes(report)
    raw = json.dumps(report, ensure_ascii=True, sort_keys=True, indent=2, allow_nan=False).encode()
    if target.exists():
        receipt = json.loads((target / 'receipt.json').read_bytes())
        require(receipt['report_sha256'] == report['report_sha256'], 'EXISTING_REPORT_CONFLICT')
        for name, sha in receipt['files'].items():
            require(name in ('report.json', 'ETF_TOP5.xlsx')
                    and hashlib.sha256((target / name).read_bytes()).hexdigest() == sha, 'EXISTING_REPORT_CONFLICT')
        require(set(receipt['files']) == {'report.json', 'ETF_TOP5.xlsx'}, 'EXISTING_REPORT_CONFLICT')
        require((target / 'report.json').read_bytes() == raw
                and (target / 'ETF_TOP5.xlsx').read_bytes() == xlsx, 'EXISTING_REPORT_CONFLICT')
        return target
    target.mkdir()
    for name, data in (('report.json', raw), ('ETF_TOP5.xlsx', xlsx)):
        with (target / name).open('xb') as handle:
            handle.write(data)
    receipt = dict(schema='etf116-simple-report-local-receipt-v1', report_sha256=report['report_sha256'],
                   files={'report.json': hashlib.sha256(raw).hexdigest(), 'ETF_TOP5.xlsx': hashlib.sha256(xlsx).hexdigest()},
                   accepted_session=False, research_only=True)
    with (target / 'receipt.json').open('xb') as handle:
        handle.write(json.dumps(receipt, sort_keys=True, indent=2).encode())
    return target


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--input-sha256', required=True)
    parser.add_argument('--frozen', type=Path, default=ROOT / 'research/etf116_v1/frozen_tickers.txt')
    output = parser.add_mutually_exclusive_group(required=True)
    output.add_argument('--output-dir', type=Path)
    output.add_argument('--plan-only', action='store_true')
    parser.add_argument('--refresh-holdings', action='store_true',
                        help='Existing cadence owner requests refresh; plan-only, GET-only.')
    args = parser.parse_args()
    try:
        require(not args.input.is_symlink(), 'INPUT_REGULAR_FILE')
        flags = os.O_RDONLY | getattr(os, 'O_NONBLOCK', 0) | getattr(os, 'O_NOFOLLOW', 0) | getattr(os, 'O_BINARY', 0)
        with os.fdopen(os.open(args.input, flags), 'rb') as handle:
            info = os.fstat(handle.fileno())
            require(stat.S_ISREG(info.st_mode) and info.st_size <= 16 * 1024 * 1024, 'INPUT_REGULAR_FILE_OR_SIZE')
            raw = handle.read(16 * 1024 * 1024 + 1)
        require(len(raw) <= 16 * 1024 * 1024, 'INPUT_SIZE')
        require(SHA.fullmatch(args.input_sha256) and hashlib.sha256(raw).hexdigest() == args.input_sha256, 'INPUT_BYTES_HASH')
        def pairs(items):
            result = {}
            for key, value in items:
                require(key not in result, 'DUPLICATE_JSON_KEY')
                result[key] = value
            return result
        def nonfinite(_):
            raise ReportRefused('NONFINITE_JSON')
        packet = json.loads(raw, object_pairs_hook=pairs, parse_constant=nonfinite)
        frozen = args.frozen.read_text(encoding='utf-8').splitlines()
        require(not args.refresh_holdings or args.plan_only, 'PLAN_ONLY_REFRESH_FLAG')
        if args.plan_only:
            print(json.dumps(plan_inputs(packet, frozen, holdings_refresh_due=args.refresh_holdings),
                             sort_keys=True, allow_nan=False))
            return 0
        report = build(packet, frozen)
        publish_local(report, args.output_dir)
        print(json.dumps(dict(state='RESEARCH_FILE_EXPORTED', session=report['session'],
                              report_sha256=report['report_sha256'], accepted_session=False)))
        return 0
    except Exception:
        # No raw input, credential URL, arbitrary exception or provider text.
        print('{"state":"BLOCKED","reason":"REPORT_INPUT_OR_EXPORT_REFUSED"}')
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
