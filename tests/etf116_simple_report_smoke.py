"""Native report regressions. All financial values here are CONTRACT_FIXTUREs."""
from __future__ import annotations

import copy
import hashlib
import io
import json
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from tools import etf116_simple_report as report

FROZEN = (ROOT / 'research/etf116_v1/frozen_tickers.txt').read_text().splitlines()
TICKERS = report.validate_universe(FROZEN)
NOW = datetime(2026, 10, 23, 1, tzinfo=timezone.utc)


def packet():
    metadata, results, holdings = {}, {}, {}
    for n, ticker in enumerate(TICKERS):
        metadata[ticker] = dict(asset_class='EQUITY' if n < 7 else 'BOND',
                                **{flag: False for flag in report.FLAGS}, currency='USD',
                                fund_name='CONTRACT_FIXTURE ' + ticker, expense_ratio=.001 + n * .00001,
                                asof='2026-10-22', published_at='2026-10-22T22:00:00Z',
                                source_url='https://example.invalid/fund/' + ticker, source_sha256='a'*64)
        windows = {h: dict(status='RESEARCH_DAILY_TR_COMPLETE', basis='MARKET_TR_INDEX',
                           anchor_session=anchor, cagr=.05 + n * .01,
                           risk=dict(risk_free_status='RESEARCH_RF_AVAILABLE', sharpe=.8 + n * .1,
                                     mdd=-.3 + n * .02, ann_volatility=.3 - n * .01))
                   for h, anchor in zip(report.HORIZONS, ('2025-10-22', '2023-10-20', '2021-10-22'))}
        results[ticker] = dict(ticker=ticker, asof='2026-10-22', series_basis='MARKET_TR_INDEX', horizons=windows)
        if n < 7:
            holdings[ticker] = [dict(rank=i, symbol='FIX' + str(i), company_name='CONTRACT_FIXTURE STOCK ' + str(i),
                                     asset_class='EQUITY', weight_pct=11 - i,
                                     asof='2026-10-22', published_at='2026-10-22T22:00:00Z',
                                     source_url='https://example.invalid/holdings/' + ticker, source_sha256='b'*64)
                                for i in range(1, 11)]
    metrics = dict(schema='etf116-metrics-candidate-output-v1', session='2026-10-22',
                   universe_sha256=report.EXPECTED_UNIVERSE_SHA256,
                   research_status='PREPARED_ONLY_NO_AUTHORITY', economic_authority=False,
                   allowed_for_sheet_publisher=False, source_input_bytes_sha256='c'*64, ticker_results=results)
    metrics['candidate_receipt_sha256'] = report.canonical_hash(metrics)
    return dict(schema=report.SCHEMA, decision_at='2026-10-22T23:00:00Z',
                metrics_candidate=metrics, fund_metadata=metadata, holdings=holdings)


def reseal(value):
    metrics = value['metrics_candidate']
    metrics.pop('candidate_receipt_sha256', None)
    metrics['candidate_receipt_sha256'] = report.canonical_hash(metrics)


class SimpleReportTests(unittest.TestCase):
    def setUp(self):
        self.value = packet()

    def build(self, value=None):
        return report.build(self.value if value is None else value, FROZEN, evaluation_time=NOW)

    def test_three_views_and_no_authority_or_return_recalculation(self):
        with patch('tools.etf116_research_metrics_h1.compute_horizons', side_effect=RuntimeError('NO_DUPLICATE_CALCULATOR')):
            result = self.build()
        self.assertEqual(len(result['top5']), 5)
        self.assertEqual(len(result['holdings']), 50)
        self.assertEqual(len(result['overlap']), 10)
        self.assertTrue(all(r['fund_count'] == 5 for r in result['overlap']))
        self.assertFalse(result['economic_authority'])
        self.assertFalse(result['accepted_session'])
        self.assertFalse(result['source_admission'])
        self.assertEqual(result['qualified_count'], 7)
        self.assertEqual(result['frozen_count'], 116)
        self.assertAlmostEqual(sum(report.WEIGHTS.values()), 1.)
        self.assertAlmostEqual(sum(report.WEIGHTS[k] for k in ('cagr_1y', 'cagr_3y', 'cagr_5y')), .6)

    def test_all_three_return_horizons_influence_ranking(self):
        baseline = self.build()
        last = baseline['top5'][-1]['ticker']
        for horizon in report.HORIZONS:
            changed = copy.deepcopy(self.value)
            changed['metrics_candidate']['ticker_results'][last]['horizons'][horizon]['cagr'] = 10.
            reseal(changed)
            before = next(r['score'] for r in baseline['top5'] if r['ticker'] == last)
            after = next(r['score'] for r in self.build(changed)['top5'] if r['ticker'] == last)
            self.assertGreater(after, before)

    def test_lower_fee_volatility_and_drawdown_are_better(self):
        for key in ('expense_ratio', 'ann_volatility', 'mdd'):
            changed = copy.deepcopy(self.value)
            ticker = TICKERS[0]
            if key == 'expense_ratio':
                changed['fund_metadata'][ticker][key] = .04
            else:
                changed['metrics_candidate']['ticker_results'][ticker]['horizons']['36']['risk'][key] = .9 if key == 'ann_volatility' else -.99
                reseal(changed)
            # Direct order-preserving normalized values, including tied cohorts.
            self.assertGreater(report.percentile([-.5, -.1, -.1], -.1), report.percentile([-.5, -.1, -.1], -.5))
            self.assertNotIn(ticker, [r['ticker'] for r in self.build(changed)['top5']])

    def test_ties_are_stable_and_share_classes_remain_distinct(self):
        for ticker in TICKERS[:7]:
            for h in report.HORIZONS:
                self.value['metrics_candidate']['ticker_results'][ticker]['horizons'][h] = copy.deepcopy(self.value['metrics_candidate']['ticker_results'][TICKERS[0]]['horizons'][h])
            self.value['fund_metadata'][ticker]['expense_ratio'] = .001
        reseal(self.value)
        self.assertEqual([r['ticker'] for r in self.build()['top5']], TICKERS[:5])
        self.assertEqual(report.percentile([1, 1, 1], 1), .5)
        for recs in self.value['holdings'].values():
            recs[0]['symbol'], recs[1]['symbol'] = 'GOOG', 'GOOGL'
        result = self.build()
        self.assertIn('GOOG', [r['symbol'] for r in result['overlap']])
        self.assertIn('GOOGL', [r['symbol'] for r in result['overlap']])

    def test_excluded_products_and_missing_history_or_rf_do_not_get_zero_scores(self):
        first = TICKERS[0]
        for flag in report.FLAGS:
            changed = copy.deepcopy(self.value)
            changed['fund_metadata'][first][flag] = True
            self.assertEqual(self.build(changed)['excluded'][first], 'OUTSIDE_GENERAL_USD_EQUITY_COHORT')
        for mode in ('history', 'rf'):
            changed = copy.deepcopy(self.value)
            window = changed['metrics_candidate']['ticker_results'][first]['horizons']['60' if mode == 'history' else '36']
            if mode == 'history':
                window['status'] = 'BLOCKED_INSUFFICIENT_HISTORY'
            else:
                window['risk']['sharpe'] = None
            reseal(changed)
            result = self.build(changed)
            self.assertEqual(result['qualified_count'], 6)
            self.assertIn(first, result['excluded'])

    def test_missing_metadata_seal_basis_or_windows_refuses(self):
        for mode in ('metadata', 'seal', 'basis', 'anchor'):
            changed = copy.deepcopy(self.value)
            if mode == 'metadata':
                del changed['fund_metadata'][TICKERS[-1]]
            elif mode == 'seal':
                changed['metrics_candidate']['ticker_results'][TICKERS[0]]['asof'] = '2026-10-21'
            else:
                window = changed['metrics_candidate']['ticker_results'][TICKERS[0]]['horizons']['12']
                window['basis' if mode == 'basis' else 'anchor_session'] = 'NAV_TR_INDEX' if mode == 'basis' else '2025-10-21'
                reseal(changed)
            with self.assertRaises(report.ReportRefused):
                self.build(changed)

    def test_nonfinite_boolean_negative_fees_refuse(self):
        for value in (float('nan'), float('inf'), True, -.001, 10):
            changed = copy.deepcopy(self.value)
            changed['fund_metadata'][TICKERS[0]]['expense_ratio'] = value
            with self.assertRaises(report.ReportRefused):
                self.build(changed)

    def test_stale_future_holdings_and_credential_urls_refuse(self):
        chosen = self.build()['top5'][0]['ticker']
        for key, value in (('asof', '2026-09-01'), ('asof', '2026-10-23'),
                           ('published_at', '2026-10-23T22:00:00Z'),
                           ('source_url', 'https://user:REAL_SECRET@example.invalid/data'),
                           ('source_url', 'https://example.invalid/data?token=REAL_SECRET')):
            changed = copy.deepcopy(self.value)
            changed['holdings'][chosen][0][key] = value
            with self.assertRaises(report.ReportRefused):
                self.build(changed)

    def test_incomplete_mixed_or_excess_holdings_refuse(self):
        chosen = self.build()['top5'][0]['ticker']
        for mode in ('missing', 'duplicate', 'mixed', 'excess'):
            changed = copy.deepcopy(self.value)
            rows = changed['holdings'][chosen]
            if mode == 'missing':
                rows.pop()
            elif mode == 'duplicate':
                rows[1]['symbol'] = rows[0]['symbol']
            elif mode == 'mixed':
                rows[0]['source_sha256'] = 'd'*64
            else:
                for row in rows:
                    row['weight_pct'] = 20
            with self.assertRaises(report.ReportRefused):
                self.build(changed)

    def test_fewer_than_five_refuses_without_export(self):
        for ticker in TICKERS[:3]:
            self.value['fund_metadata'][ticker]['inverse'] = True
        with self.assertRaisesRegex(report.ReportRefused, 'FEWER_THAN_FIVE'):
            self.build()

    def test_workbook_has_three_tabs_literal_text_and_numeric_units(self):
        from openpyxl import load_workbook
        chosen = self.build()['top5'][0]['ticker']
        self.value['fund_metadata'][chosen]['fund_name'] = '=HYPERLINK("https://example.invalid","FIXTURE")'
        book = load_workbook(io.BytesIO(report.workbook_bytes(self.build())), data_only=False)
        self.assertEqual(book.sheetnames, ['ETF TOP5', 'TOP10 HOLDINGS', 'OVERLAP'])
        self.assertEqual(book['ETF TOP5']['C5'].data_type, 's')
        self.assertEqual(book['ETF TOP5']['K5'].number_format, '0.00%;[Red](0.00%);"–"')
        self.assertEqual(book['TOP10 HOLDINGS'].max_row, 54)
        self.assertFalse(any(c.data_type == 'f' for sheet in book for row in sheet for c in row))

    def test_exact_local_reuse_and_tamper_never_overwrite(self):
        result = self.build()
        with tempfile.TemporaryDirectory() as td:
            path = report.publish_local(result, td)
            before = (path / 'ETF_TOP5.xlsx').read_bytes()
            self.assertEqual(report.publish_local(result, td), path)
            self.assertEqual(report.workbook_bytes(result), before)
            (path / 'report.json').write_bytes(b'tampered')
            with self.assertRaises(report.ReportRefused):
                report.publish_local(result, td)
            self.assertEqual((path / 'ETF_TOP5.xlsx').read_bytes(), before)

    def test_rehashed_cache_cannot_replace_display_values(self):
        result = self.build()
        with tempfile.TemporaryDirectory() as td:
            path = report.publish_local(result, td)
            receipt = json.loads((path / 'receipt.json').read_bytes())
            (path / 'ETF_TOP5.xlsx').write_bytes(b'FORGED_WORKBOOK')
            receipt['files']['ETF_TOP5.xlsx'] = hashlib.sha256(b'FORGED_WORKBOOK').hexdigest()
            (path / 'receipt.json').write_text(json.dumps(receipt))
            with self.assertRaises(report.ReportRefused):
                report.publish_local(result, td)

    def test_cli_hash_failure_prints_no_secret_and_no_output(self):
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            source = base / 'input.json'
            source.write_bytes(b'{"REAL_SECRET":"https://user:token@example.invalid"}')
            command = [sys.executable, '-O' if sys.flags.optimize else '-B', str(ROOT / 'tools/etf116_simple_report.py'),
                       '--input', str(source), '--input-sha256', '0'*64, '--output-dir', str(base / 'output')]
            child = subprocess.run(command, capture_output=True, text=True)
            self.assertEqual(child.returncode, 2)
            self.assertNotIn('REAL_SECRET', child.stdout + child.stderr)
            self.assertNotIn('https://', child.stdout + child.stderr)
            self.assertFalse((base / 'output').exists())


if __name__ == '__main__':
    unittest.main()
